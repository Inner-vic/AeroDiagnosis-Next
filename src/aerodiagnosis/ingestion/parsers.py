"""Deterministic, offline parsers for the first v3 ingestion slice."""

from __future__ import annotations

import csv
import io
from collections.abc import Mapping, Sequence
from dataclasses import dataclass
from pathlib import Path
from typing import Protocol

from pypdf import PdfReader

from aerodiagnosis.domain import EvidenceLocator


class DocumentParseError(ValueError):
    pass


@dataclass(frozen=True, slots=True)
class ParsedChunk:
    content: str
    locator: EvidenceLocator


@dataclass(frozen=True, slots=True)
class ParsedDocument:
    parser_version: str
    media_type: str
    canonical_content: bytes
    chunks: tuple[ParsedChunk, ...]


class DocumentParser(Protocol):
    @property
    def parser_version(self) -> str: ...

    def parse(self, content: bytes) -> ParsedDocument: ...


def _decode_text(content: bytes) -> str:
    if not content:
        raise DocumentParseError("document is empty")
    encodings = (
        ("utf-8-sig", "utf-16")
        if content.startswith((b"\xff\xfe", b"\xfe\xff"))
        else ("utf-8-sig",)
    )
    for encoding in encodings:
        try:
            return content.decode(encoding)
        except UnicodeDecodeError:
            continue
    raise DocumentParseError("document must be valid UTF-8 or BOM-marked UTF-16")


def _canonical_text(content: bytes) -> str:
    text = _decode_text(content).replace("\r\n", "\n").replace("\r", "\n")
    if not text.strip():
        raise DocumentParseError("document contains no text")
    return text


class PlainTextParser:
    parser_version = "plain-text@1"

    def __init__(self, *, chunk_size: int = 800, overlap: int = 100) -> None:
        if chunk_size < 100 or overlap < 0 or overlap >= chunk_size:
            raise ValueError("chunk_size must be at least 100 and overlap must be smaller")
        self._chunk_size = chunk_size
        self._step = chunk_size - overlap

    def parse(self, content: bytes) -> ParsedDocument:
        text = _canonical_text(content)
        chunks = []
        for start in range(0, len(text), self._step):
            end = min(start + self._chunk_size, len(text))
            value = text[start:end].strip()
            if value:
                chunks.append(
                    ParsedChunk(
                        content=value,
                        locator=EvidenceLocator(
                            kind="character_range",
                            coordinates={"start": start, "end": end},
                        ),
                    )
                )
            if end == len(text):
                break
        return ParsedDocument(
            parser_version=self.parser_version,
            media_type="text/plain",
            canonical_content=text.encode("utf-8"),
            chunks=tuple(chunks),
        )


class CsvParser:
    parser_version = "csv@1"

    def __init__(self, *, rows_per_chunk: int = 20) -> None:
        if rows_per_chunk < 1:
            raise ValueError("rows_per_chunk must be positive")
        self._rows_per_chunk = rows_per_chunk

    def parse(self, content: bytes) -> ParsedDocument:
        text = _canonical_text(content)
        try:
            rows = list(csv.reader(io.StringIO(text, newline=""), strict=True))
        except csv.Error as exc:
            raise DocumentParseError("invalid CSV content") from exc
        if not rows or not any(cell.strip() for row in rows for cell in row):
            raise DocumentParseError("CSV contains no cells")
        header = rows[0]
        data_rows = rows[1:]
        chunks = []
        if not data_rows:
            chunks.append(
                ParsedChunk(
                    content="Columns: " + " | ".join(header),
                    locator=EvidenceLocator(
                        kind="csv_rows",
                        coordinates={"row_start": 1, "row_end": 1},
                    ),
                )
            )
        for offset in range(0, len(data_rows), self._rows_per_chunk):
            batch = data_rows[offset : offset + self._rows_per_chunk]
            lines = ["Columns: " + " | ".join(header)]
            for index, row in enumerate(batch, start=offset + 2):
                fields = []
                for column, value in enumerate(row):
                    label = (
                        header[column]
                        if column < len(header) and header[column]
                        else f"column_{column + 1}"
                    )
                    fields.append(f"{label}: {value}")
                lines.append(f"Row {index}: " + " | ".join(fields))
            chunks.append(
                ParsedChunk(
                    content="\n".join(lines),
                    locator=EvidenceLocator(
                        kind="csv_rows",
                        coordinates={
                            "row_start": offset + 2,
                            "row_end": offset + len(batch) + 1,
                        },
                    ),
                )
            )
        return ParsedDocument(
            parser_version=self.parser_version,
            media_type="text/csv",
            canonical_content=text.encode("utf-8"),
            chunks=tuple(chunks),
        )


class PdfParser:
    parser_version = "pdf@1"

    def parse(self, content: bytes) -> ParsedDocument:
        try:
            reader = PdfReader(io.BytesIO(content))
        except Exception as exc:
            raise DocumentParseError("invalid PDF content") from exc
        chunks: list[ParsedChunk] = []
        texts: list[str] = []
        for page_number, page in enumerate(reader.pages, start=1):
            text = (page.extract_text() or "").strip()
            if not text:
                continue
            texts.append(text)
            chunks.append(
                ParsedChunk(
                    content=text,
                    locator=EvidenceLocator(
                        kind="pdf_page",
                        coordinates={"page": page_number},
                    ),
                )
            )
        if not chunks:
            raise DocumentParseError("PDF contains no extractable text")
        return ParsedDocument(
            parser_version=self.parser_version,
            media_type="application/pdf",
            canonical_content="\n\n".join(texts).encode("utf-8"),
            chunks=tuple(chunks),
        )


class DocxParser:
    parser_version = "docx@1"

    def parse(self, content: bytes) -> ParsedDocument:
        from docx import Document as DocxDocument

        try:
            document = DocxDocument(io.BytesIO(content))
        except Exception as exc:
            raise DocumentParseError("invalid DOCX content") from exc
        parts: list[str] = [
            paragraph.text.strip()
            for paragraph in document.paragraphs
            if paragraph.text.strip()
        ]
        for table in document.tables:
            for row in table.rows:
                values = [cell.text.strip() for cell in row.cells]
                if any(values):
                    parts.append("\t".join(values))
        canonical = "\n".join(parts)
        if not canonical.strip():
            raise DocumentParseError("DOCX contains no text")
        return ParsedDocument(
            parser_version=self.parser_version,
            media_type=(
                "application/vnd.openxmlformats-officedocument."
                "wordprocessingml.document"
            ),
            canonical_content=canonical.encode("utf-8"),
            chunks=(
                ParsedChunk(
                    content=canonical,
                    locator=EvidenceLocator(
                        kind="docx_document",
                        coordinates={"parts": len(parts)},
                    ),
                ),
            ),
        )


class XlsxParser:
    parser_version = "xlsx@1"

    def parse(self, content: bytes) -> ParsedDocument:
        from openpyxl import load_workbook

        try:
            workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        except Exception as exc:
            raise DocumentParseError("invalid XLSX content") from exc
        chunks: list[ParsedChunk] = []
        sheet_texts: list[str] = []
        for worksheet in workbook.worksheets:
            lines: list[str] = []
            for row in worksheet.iter_rows(values_only=True):
                if any(value is not None and str(value).strip() for value in row):
                    lines.append(
                        " | ".join("" if value is None else str(value) for value in row)
                    )
            text = "\n".join(lines)
            if text.strip():
                chunks.append(
                    ParsedChunk(
                        content=text,
                        locator=EvidenceLocator(
                            kind="xlsx_sheet",
                            coordinates={
                                "sheet": worksheet.title,
                                "rows": len(lines),
                            },
                        ),
                    )
                )
                sheet_texts.append(text)
        if not chunks:
            raise DocumentParseError("XLSX contains no rows")
        return ParsedDocument(
            parser_version=self.parser_version,
            media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            canonical_content="\n\n".join(sheet_texts).encode("utf-8"),
            chunks=tuple(chunks),
        )


class HtmlParser:
    parser_version = "html@1"

    def parse(self, content: bytes) -> ParsedDocument:
        from bs4 import BeautifulSoup

        try:
            soup = BeautifulSoup(content, "lxml")
        except Exception as exc:
            raise DocumentParseError("invalid HTML content") from exc
        for element in soup(["script", "style", "noscript"]):
            element.decompose()
        text = "\n".join(
            line.strip()
            for line in soup.get_text("\n").splitlines()
            if line.strip()
        )
        if not text:
            raise DocumentParseError("HTML contains no text")
        return ParsedDocument(
            parser_version=self.parser_version,
            media_type="text/html",
            canonical_content=text.encode("utf-8"),
            chunks=(
                ParsedChunk(
                    content=text,
                    locator=EvidenceLocator(
                        kind="html_document",
                        coordinates={"lines": len(text.splitlines())},
                    ),
                ),
            ),
        )


class ParserRegistry:
    def __init__(self, parsers: Mapping[str, DocumentParser] | None = None) -> None:
        defaults: dict[str, DocumentParser] = {
            ".txt": PlainTextParser(),
            ".md": PlainTextParser(),
            ".markdown": PlainTextParser(),
            ".csv": CsvParser(),
            ".pdf": PdfParser(),
            ".docx": DocxParser(),
            ".xlsx": XlsxParser(),
            ".html": HtmlParser(),
            ".htm": HtmlParser(),
        }
        self._parsers = {key.lower(): value for key, value in (parsers or defaults).items()}

    @property
    def supported_extensions(self) -> Sequence[str]:
        return tuple(sorted(self._parsers))

    def parse(self, display_name: str, content: bytes) -> ParsedDocument:
        suffix = Path(display_name).suffix.lower()
        try:
            parser = self._parsers[suffix]
        except KeyError as exc:
            supported = ", ".join(self.supported_extensions)
            raise DocumentParseError(
                f"unsupported document extension {suffix or '<none>'}; supported: {supported}"
            ) from exc
        return parser.parse(content)
