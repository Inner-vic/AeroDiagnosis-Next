"""
文档解析 Agent — 多模态文档解析，支持 PDF / 图片 / 表格 / 纯文本
"""

from __future__ import annotations

import hashlib
import os
from dataclasses import dataclass, field
from enum import Enum
from typing import Any

from langchain_core.messages import HumanMessage, SystemMessage
from langchain_openai import ChatOpenAI

from config import settings


class DocType(str, Enum):
    PDF = "pdf"
    IMAGE = "image"
    TABLE = "table"
    TEXT = "text"
    MARKDOWN = "markdown"
    UNKNOWN = "unknown"


@dataclass
class DocumentChunk:
    content: str
    doc_id: str
    chunk_index: int
    doc_type: DocType
    metadata: dict[str, Any] = field(default_factory=dict)
    embedding: list[float] | None = None

    @property
    def chunk_id(self) -> str:
        return f"{self.doc_id}#chunk-{self.chunk_index}"


class DocParserAgent:
    SUPPORTED_EXTENSIONS: dict[str, DocType] = {
        ".pdf": DocType.PDF,
        ".png": DocType.IMAGE,
        ".jpg": DocType.IMAGE,
        ".jpeg": DocType.IMAGE,
        ".csv": DocType.TABLE,
        ".xlsx": DocType.TABLE,
        ".xls": DocType.TABLE,
        ".txt": DocType.TEXT,
        ".md": DocType.MARKDOWN,
        ".docx": DocType.TEXT,
        ".doc": DocType.TEXT,
        ".pptx": DocType.TEXT,
        ".html": DocType.MARKDOWN,
    }

    CHUNK_SIZE = 512
    CHUNK_OVERLAP = 64

    def __init__(self) -> None:
        self.llm = ChatOpenAI(
            model=settings.openai_model,
            api_key=settings.openai_api_key,
            base_url=settings.openai_base_url,
            temperature=0,
        )

    async def parse(self, file_path: str) -> list[DocumentChunk]:
        doc_type = self._classify(file_path)
        doc_id = self._make_doc_id(file_path)
        raw_texts: list[str] = []

        if doc_type == DocType.PDF:
            raw_texts = await self._parse_pdf(file_path)
        elif doc_type == DocType.IMAGE:
            raw_texts = await self._parse_image(file_path)
        elif doc_type == DocType.TABLE:
            raw_texts = await self._parse_table(file_path)
        elif hasattr(self, '_maybe_docx') and self._maybe_docx(file_path):
            raw_texts = self._parse_docx(file_path)
        else:
            raw_texts = self._parse_text(file_path)

        return self._chunk_texts(raw_texts, doc_id, doc_type, file_path)

    @staticmethod
    def _maybe_docx(file_path: str) -> bool:
        return os.path.splitext(file_path)[1].lower() in (".docx", ".doc", ".pptx", ".html")

    @staticmethod
    def _parse_docx(file_path: str) -> list[str]:
        try:
            import docx
            doc = docx.Document(file_path)
            paragraphs = []
            for para in doc.paragraphs:
                text = para.text.strip()
                if text:
                    paragraphs.append(text)
            if not paragraphs:
                # Try extracting tables
                for table in doc.tables:
                    for row in table.rows:
                        cells = [cell.text.strip() for cell in row.cells]
                        paragraphs.append(" | ".join(cells))
            return paragraphs if paragraphs else ["[Empty DOCX]"]
        except Exception:
            return ["[DOCX parse failed] " + file_path]

    async def parse_batch(self, file_paths: list[str]) -> list[DocumentChunk]:
        all_chunks: list[DocumentChunk] = []
        for fp in file_paths:
            all_chunks.extend(await self.parse(fp))
        return all_chunks

    def _classify(self, file_path: str) -> DocType:
        ext = os.path.splitext(file_path)[1].lower()
        return self.SUPPORTED_EXTENSIONS.get(ext, DocType.UNKNOWN)

    @staticmethod
    def _make_doc_id(file_path: str) -> str:
        return hashlib.sha256(file_path.encode()).hexdigest()[:16]

    async def _parse_pdf(self, file_path: str) -> list[str]:
        texts: list[str] = []
        try:
            from PyPDF2 import PdfReader
            reader = PdfReader(file_path)
            for page in reader.pages:
                page_text = page.extract_text() or ""
                if page_text.strip():
                    texts.append(page_text.strip())
        except Exception:
            texts.append(f"[PDF parse failed] {file_path}")
        if not texts:
            texts = await self._pdf_vision_fallback(file_path)
        return texts

    async def _pdf_vision_fallback(self, file_path: str) -> list[str]:
        try:
            from pdf2image import convert_from_path
            images = convert_from_path(file_path, dpi=150, first_page=1, last_page=5)
            texts: list[str] = []
            for img in images:
                desc = await self._describe_image_with_llm(img)
                texts.append(desc)
            return texts
        except Exception:
            return [f"[PDF vision fallback failed] {file_path}"]

    async def _parse_image(self, file_path: str) -> list[str]:
        texts: list[str] = []
        ocr_text = self._ocr(file_path)
        if ocr_text.strip():
            texts.append(ocr_text)
        from PIL import Image
        img = Image.open(file_path)
        desc = await self._describe_image_with_llm(img)
        texts.append(desc)
        return texts

    @staticmethod
    def _ocr(file_path: str) -> str:
        try:
            import pytesseract
            from PIL import Image
            return pytesseract.image_to_string(Image.open(file_path), lang="chi_sim+eng")
        except Exception:
            return ""

    async def _describe_image_with_llm(self, image: Any) -> str:
        import base64
        import io
        buf = io.BytesIO()
        image.save(buf, format="PNG")
        b64 = base64.b64encode(buf.getvalue()).decode()
        messages = [
            SystemMessage(content="你是一个专业的文档分析助手，请详细描述图片中的内容，包括文字、表格、图表信息。"),
            HumanMessage(content=[
                {"type": "text", "text": "请描述这张图片的所有内容："},
                {"type": "image_url", "image_url": {"url": f"data:image/png;base64,{b64}"}},
            ]),
        ]
        resp = await self.llm.ainvoke(messages)
        return resp.content

    async def _parse_table(self, file_path: str) -> list[str]:
        ext = os.path.splitext(file_path)[1].lower()
        try:
            if ext == ".csv":
                return self._parse_csv(file_path)
            else:
                return self._parse_excel(file_path)
        except Exception:
            return [f"[Table parse failed] {file_path}"]

    @staticmethod
    def _parse_csv(file_path: str) -> list[str]:
        import csv
        texts: list[str] = []
        with open(file_path, encoding="utf-8") as f:
            reader = csv.DictReader(f)
            headers = reader.fieldnames or []
            rows: list[str] = []
            for row in reader:
                rows.append(" | ".join(f"{h}: {row.get(h, '')}" for h in headers))
            for i in range(0, len(rows), 20):
                batch = rows[i:i + 20]
                texts.append(f"Headers: {' | '.join(headers)}\n" + "\n".join(batch))
        return texts or ["[Empty CSV]"]

    @staticmethod
    def _parse_excel(file_path: str) -> list[str]:
        try:
            import openpyxl
            wb = openpyxl.load_workbook(file_path, read_only=True)
            texts: list[str] = []
            for sheet in wb.worksheets:
                rows = list(sheet.iter_rows(values_only=True))
                if not rows:
                    continue
                headers = [str(c) if c else "" for c in rows[0]]
                data_rows: list[str] = []
                for row in rows[1:]:
                    data_rows.append(" | ".join(
                        f"{headers[j]}: {row[j]}" if j < len(headers) else str(row[j])
                        for j in range(len(row))
                    ))
                for i in range(0, len(data_rows), 20):
                    batch = data_rows[i:i + 20]
                    texts.append(f"Sheet: {sheet.title}\nHeaders: {' | '.join(headers)}\n" + "\n".join(batch))
            return texts or ["[Empty Excel]"]
        except Exception:
            return [f"[Excel parse failed] {file_path}"]

    @staticmethod
    def _read_text_file(file_path: str) -> str:
        """读取文本文件，自动尝试多种编码"""
        for enc in ("utf-8", "utf-16", "gbk", "gb2312", "gb18030", "latin-1"):
            try:
                with open(file_path, encoding=enc) as f:
                    return f.read()
            except (UnicodeDecodeError, UnicodeError):
                continue
        # Last resort — ignore errors
        with open(file_path, encoding="utf-8", errors="replace") as f:
            return f.read()

    @staticmethod
    def _parse_text(file_path: str) -> list[str]:
        return [DocParserAgent._read_text_file(file_path)]

    def _chunk_texts(self, texts: list[str], doc_id: str, doc_type: DocType, source: str) -> list[DocumentChunk]:
        chunks: list[DocumentChunk] = []
        idx = 0
        for text in texts:
            start = 0
            while start < len(text):
                end = start + self.CHUNK_SIZE
                content = text[start:end]
                if content.strip():
                    chunks.append(DocumentChunk(
                        content=content.strip(),
                        doc_id=doc_id,
                        chunk_index=idx,
                        doc_type=doc_type,
                        metadata={"source": source, "char_start": start, "char_end": end},
                    ))
                    idx += 1
                start = end - self.CHUNK_OVERLAP
        return chunks
