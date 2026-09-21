from __future__ import annotations

import io

import pytest
from docx import Document as DocxDocument
from openpyxl import Workbook

from aerodiagnosis.ingestion import DocumentParseError, ParserRegistry
from aerodiagnosis.ingestion.parsers import DocxParser, HtmlParser, PdfParser, XlsxParser


def test_pdf_parser_extracts_page_scoped_chunks(monkeypatch: object) -> None:
    class FakePage:
        def extract_text(self) -> str:
            return "Compressor stall guidance.\nCheck EGT trend."

    class FakeReader:
        def __init__(self, _stream: object) -> None:
            self.pages = [FakePage()]

    monkeypatch.setattr("aerodiagnosis.ingestion.parsers.PdfReader", FakeReader)

    parsed = PdfParser().parse(b"fake-pdf")

    assert parsed.media_type == "application/pdf"
    assert parsed.parser_version == "pdf@1"
    assert parsed.chunks[0].locator.kind == "pdf_page"
    assert parsed.chunks[0].locator.coordinates["page"] == 1


def test_docx_parser_preserves_paragraph_text_and_tables() -> None:
    document = DocxDocument()
    document.add_paragraph("Compressor inspection")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "EGT"
    table.cell(0, 1).text = "710"
    stream = io.BytesIO()
    document.save(stream)

    parsed = DocxParser().parse(stream.getvalue())

    assert parsed.media_type == (
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    )
    assert "Compressor inspection" in parsed.canonical_content.decode()
    assert "EGT" in parsed.canonical_content.decode()


def test_xlsx_parser_preserves_sheet_and_row_data() -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "EGT"
    sheet.append(["parameter", "value"])
    sheet.append(["egt", 710])
    stream = io.BytesIO()
    workbook.save(stream)

    parsed = XlsxParser().parse(stream.getvalue())

    assert parsed.media_type == (
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    assert parsed.chunks[0].locator.kind == "xlsx_sheet"
    assert parsed.chunks[0].locator.coordinates["sheet"] == "EGT"
    assert "egt" in parsed.canonical_content.decode()


def test_html_parser_removes_script_content() -> None:
    content = (
        b"<html><body><h1>EGT guidance</h1>"
        b"<script>alert('unsafe')</script>"
        b"<p>Check compressor trend.</p></body></html>"
    )

    parsed = HtmlParser().parse(content)

    assert parsed.media_type == "text/html"
    assert "unsafe" not in parsed.canonical_content.decode()
    assert "EGT guidance" in parsed.canonical_content.decode()
    assert "Check compressor trend" in parsed.canonical_content.decode()


def test_registry_advertises_structured_documents() -> None:
    registry = ParserRegistry()

    assert ".docx" in registry.supported_extensions
    assert ".html" in registry.supported_extensions
    assert ".pdf" in registry.supported_extensions
    assert ".xlsx" in registry.supported_extensions
    with pytest.raises(DocumentParseError, match="unsupported document extension"):
        registry.parse("unsupported.xyz", b"content")
